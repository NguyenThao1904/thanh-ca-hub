"""The library as an Excel file, with the same search, filters and order as on screen."""

from __future__ import annotations

import io
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from ..schemas import SongListItem, SortOption
from .search import search_songs

BATCH = 1000

# Column titles in the app's two languages (same words as on screen).
LABELS = {
    "vi": {
        "sheet": "Thư viện bài hát",
        "columns": [
            "Vị trí",
            "Số bài",
            "Tên bài",
            "Tác giả",
            "Câu đầu",
            "Danh mục",
            "Tập",
            "Trang / vị trí",
            "PDF",
            "Ảnh bản nhạc",
            "Ngày thêm",
            "Cập nhật lần cuối",
        ],
    },
    "en": {
        "sheet": "Song Library",
        "columns": [
            "Location",
            "Song no.",
            "Title",
            "Composer",
            "First sentence",
            "Category",
            "Binder",
            "Page / position",
            "PDF",
            "Scanned pages",
            "Date added",
            "Last updated",
        ],
    },
}
WIDTHS = [11, 9, 36, 24, 50, 16, 7, 9, 7, 9, 12, 12]


def all_matching_songs(
    db: Session, *, query: str, category_id: int | None, binder: int | None, sort: SortOption | None
) -> list[SongListItem]:
    songs: list[SongListItem] = []
    page = 1
    while True:
        rows, _total = search_songs(
            db, query=query, category_id=category_id, binder=binder, sort=sort, page=page, page_size=BATCH
        )
        songs.extend(rows)
        if len(rows) < BATCH:
            return songs
        page += 1


def build_workbook(songs: list[SongListItem], locale: str, zone: ZoneInfo) -> bytes:
    labels = LABELS.get(locale, LABELS["vi"])

    def day(value: datetime) -> str:
        # yyyy-mm-dd in the choir's time zone: sorts and filters well in a spreadsheet.
        return value.astimezone(zone).strftime("%Y-%m-%d")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = labels["sheet"]
    sheet.append(labels["columns"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for song in songs:
        sheet.append(
            [
                song.location,
                song.song_number,
                song.title,
                song.composer,
                song.first_sentence,
                song.category_name,
                song.binder_number,
                song.page_number,
                "PDF" if song.has_pdf else "",
                song.image_count or None,
                day(song.created_at),
                day(song.updated_at),
            ]
        )
    for index, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    # Keep the header visible while scrolling.
    sheet.freeze_panes = "A2"

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
