"""Excel export, serving the web app, the database schema, the command line and S3 storage."""

from __future__ import annotations

import io
import sqlite3
import zipfile
from pathlib import Path

import boto3
from alembic.autogenerate import compare_metadata
from alembic.migration import MigrationContext
from fastapi.testclient import TestClient
from moto import mock_aws
from openpyxl import load_workbook

from thanhca.app import create_app
from thanhca.cli import main as cli
from thanhca.config import Settings
from thanhca.models import Base
from thanhca.storage import S3Storage

from helpers import category_id, signed_in


def workbook(response):
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert response.headers["content-disposition"].startswith('attachment; filename="songs-')
    return load_workbook(io.BytesIO(response.content)).active


def test_export_follows_the_current_search_and_language(app, member, sample_songs):
    sheet = workbook(member.get("/api/export/songs", params={"q": "dang loi", "lang": "en"}))
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:4] == ("Location", "Song no.", "Title", "Composer")
    assert rows[1][:3] == ("DL-1.03", 125, "Xin Dâng Lời Cảm Tạ")
    assert len(rows) - 1 == member.get("/api/songs", params={"q": "dang loi"}).json()["total"]
    assert sheet.freeze_panes == "A2"


def test_export_all_songs_in_shelf_order_in_vietnamese(app, member, sample_songs):
    sheet = workbook(member.get("/api/export/songs", params={"category": category_id(app, "NL")}))
    rows = list(sheet.iter_rows(values_only=True))
    assert sheet.title == "Thư viện bài hát"
    assert rows[0][0] == "Vị trí"
    assert [row[0] for row in rows[1:]] == ["NL-1.01", "NL-1.02", "NL-1.03", "NL-1.10", "NL-1.11"]
    assert rows[1][10].count("-") == 2  # yyyy-mm-dd


def test_migrations_match_the_models(app):
    with app.state.db.engine.connect() as connection:
        differences = compare_metadata(MigrationContext.configure(connection), Base.metadata)
    assert differences == []


def test_serves_the_web_app_with_a_strict_content_security_policy(tmp_path, settings):
    dist = tmp_path / "dist"
    (dist / "assets").mkdir(parents=True)
    (dist / "index.html").write_text("<!doctype html><title>app</title>", encoding="utf-8")
    (dist / "assets" / "index-abc123.js").write_text("console.log(1)", encoding="utf-8")
    (dist / "robots.txt").write_text("User-agent: *\nDisallow: /", encoding="utf-8")
    client = TestClient(create_app(settings.model_copy(update={"frontend_dist": dist})))

    page = client.get("/songs/12/edit")
    assert page.status_code == 200
    assert "<title>app</title>" in page.text
    assert page.headers["cache-control"] == "no-cache"
    assert "default-src 'self'" in page.headers["content-security-policy"]
    assert "frame-ancestors 'self'" in page.headers["content-security-policy"]

    asset = client.get("/assets/index-abc123.js")
    assert asset.headers["cache-control"] == "public, max-age=31536000, immutable"
    assert client.get("/robots.txt").text.startswith("User-agent")
    assert client.head("/songs").status_code == 200
    assert client.get("/../backend/pyproject.toml").text.startswith("<!doctype html>")
    unknown_api = client.get("/api/nothing-here")
    assert unknown_api.status_code == 404
    assert unknown_api.json() == {"error": {"message": "errors.notFound"}}


def test_api_documentation_is_available(anonymous):
    assert anonymous.get("/api/openapi.json").json()["info"]["title"] == "Thánh Ca Hub API"


def test_cli_creates_and_promotes_admins_and_loads_samples(tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    monkeypatch.setenv("DATABASE_URL", "")
    monkeypatch.setenv("STORAGE_BACKEND", "local")
    assert (
        cli(
            [
                "create-admin",
                "--email",
                "Choir.Admin@Example.org",
                "--password",
                "a-long-password",
                "--name",
                "Ca Trưởng",
            ]
        )
        == 0
    )
    assert "Created admin account choir.admin@example.org" in capsys.readouterr().out
    assert cli(["create-admin", "--email", "choir.admin@example.org", "--password", "another-password"]) == 0
    assert "Updated admin account" in capsys.readouterr().out
    assert cli(["create-admin", "--email", "x@example.org", "--password", "short"]) == 1
    # A placeholder copied from the instructions is not an email address.
    assert cli(["create-admin", "--email", "…", "--password", "a-long-password", "--name", "…"]) == 1
    assert "Not a valid email address" in capsys.readouterr().err
    assert cli(["seed"]) == 0
    assert "Added 36 sample songs" in capsys.readouterr().out
    assert cli(["seed"]) == 0
    assert "Added 0 sample songs" in capsys.readouterr().out
    assert cli(["cleanup"]) == 0
    assert (tmp_path / "thanhca.db").is_file()

    (tmp_path / "files" / "songs").mkdir(parents=True)
    (tmp_path / "files" / "songs" / ("b" * 32 + ".pdf")).write_bytes(b"%PDF-1.4")
    archive = tmp_path / "backup.zip"
    assert cli(["backup", "--output", str(archive)]) == 0
    assert "database and 1 files" in capsys.readouterr().out
    with zipfile.ZipFile(archive) as backup:
        assert sorted(backup.namelist()) == ["files/songs/" + "b" * 32 + ".pdf", "thanhca.db"]
        restored = tmp_path / "restored.db"
        restored.write_bytes(backup.read("thanhca.db"))
    connection = sqlite3.connect(restored)
    try:
        assert connection.execute("select count(*) from songs").fetchone() == (36,)
    finally:
        connection.close()


def test_settings_accept_hosted_postgres_urls():
    settings = Settings(_env_file=None, database_url="postgres://user:secret@db.example.org:5432/choir")
    assert settings.sqlalchemy_url == "postgresql+psycopg://user:secret@db.example.org:5432/choir"
    assert Settings(_env_file=None, time_zone="Mars/Olympus").time_zone == "UTC"
    relative = Settings(_env_file=None, database_url="sqlite:///data/test.db").sqlalchemy_url
    assert Path(relative.removeprefix("sqlite:///")).is_absolute()


@mock_aws
def test_s3_storage_keeps_files_private_and_serves_expiring_links(app, settings, tmp_path):
    boto3.client("s3", region_name="us-east-1").create_bucket(Bucket="choir-files")
    s3_settings = settings.model_copy(
        update={
            "storage_backend": "s3",
            "s3_bucket": "choir-files",
            "s3_region": "us-east-1",
            "s3_access_key_id": "test",
            "s3_secret_access_key": "test",
            "s3_prefix": "thanhca",
            "frontend_dist": tmp_path,
        }
    )
    (tmp_path / "index.html").write_text("<!doctype html><title>app</title>", encoding="utf-8")
    storage = S3Storage(s3_settings)
    storage.save("songs/" + "a" * 32 + ".pdf", io.BytesIO(b"%PDF-1.4"), "application/pdf")
    assert storage.exists("songs/" + "a" * 32 + ".pdf")
    keys = [item["Key"] for item in storage.client.list_objects_v2(Bucket="choir-files")["Contents"]]
    assert keys == ["thanhca/songs/" + "a" * 32 + ".pdf"]

    response = storage.response(
        "songs/" + "a" * 32 + ".pdf", content_type="application/pdf", file_name="Bài hát.pdf", download=True
    )
    assert response.status_code == 302
    link = response.headers["location"]
    assert "X-Amz-Expires=3600" in link and "response-content-disposition=attachment" in link
    assert storage.external_origin.startswith("https://")

    storage.delete(["songs/" + "a" * 32 + ".pdf"])
    assert not storage.exists("songs/" + "a" * 32 + ".pdf")

    # The whole app works with S3: upload, attach, and open through a redirect.
    s3_app = create_app(s3_settings)
    client = signed_in(s3_app, "s3admin@choir.test", role="admin")
    category = client.post("/api/categories", json={"name": "Thử", "code": "TS", "binderCount": 1}).json()
    song = client.post(
        "/api/songs", json={"categoryId": category["id"], "binderNumber": 1, "pageNumber": 1, "title": "S3"}
    ).json()
    upload = client.post(
        "/api/uploads", data={"fileType": "pdf"}, files={"file": ("s3.pdf", b"%PDF-1.4 x", "application/pdf")}
    )
    assert upload.status_code == 201
    client.put(f"/api/songs/{song['id']}/files", json={"pdf": {"uploadId": upload.json()["id"]}})
    opened = client.get(f"/api/files/{upload.json()['id']}", follow_redirects=False)
    assert opened.status_code == 302
    assert "choir-files" in opened.headers["location"]
    # The page's Content Security Policy lets images and the PDF preview load from the bucket.
    csp = client.get("/songs").headers["content-security-policy"]
    assert f"img-src 'self' blob: data: {storage.external_origin}" in csp
    assert f"frame-src 'self' {storage.external_origin}" in csp
    s3_app.state.db.dispose()
